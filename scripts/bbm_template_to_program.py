"""Convert a Barbell Medicine-style weekly template (.xlsx) into a Strength Tracker program file.

Usage: python3 scripts/bbm_template_to_program.py Template.xlsx out.program.json  (needs openpyxl)
Then use Programs → Import in the app. Keep the output private: it contains the paid program.
"""
import json, math, re, sys
import openpyxl

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)

NAME_MAP = {
    'competition style low bar back squat (w/ belt)': 'squat',
    'competition-style bench press': 'bench-press',
    'competition-style conventional deadlift w/ belt': 'deadlift',
    'touch and go bench press': 'touch-and-go-bench',
    'front squat': 'front-squat',
    'low bar back squat, no belt': 'beltless-squat',
    'seated overhead press': 'seated-press',
    'low incline (15*) bench press': 'bbm-low-incline-bench',
    'high incline bench press, normal grip': 'incline-bench-press',
    'alternate stance deadlift': 'sumo-deadlift',
    'slingshot bench': 'slingshot-bench',
    'pendlay row': 'pendlay-row',
    'close grip floor press': 'bbm-close-grip-floor-press',
    '2 count paused bench press': 'bench-2ct-pause',
    '2 count paused squat': 'pause-squat',
    'standing overhead press': 'press',
    '2 count paused deadlift': 'paused-deadlift',
    'floor press': 'floor-press',
    'romanian deadlift': 'romanian-deadlift',
    'feet up bench press': 'feet-up-bench',
    'pin bench (chest level)': 'bbm-pin-bench-chest',
    'pin squat (below parallel)': 'pin-squat',
    'wide grip bench press': 'wide-grip-bench',
    '2" deficit deadlift': 'deficit-deadlift',
    'close grip bench press': 'close-grip-bench',
    'stiff legged deadlift off 2" deficit': 'bbm-deficit-sldl',
}
CUSTOM = [
    dict(id='bbm-low-incline-bench', name='Low Incline Bench Press (15°)', category='bench', equipment='barbell', primary=['chest', 'frontDelts'], secondary=['triceps']),
    dict(id='bbm-close-grip-floor-press', name='Close Grip Floor Press', category='bench', equipment='barbell', primary=['triceps', 'chest'], secondary=['frontDelts']),
    dict(id='bbm-pin-bench-chest', name='Pin Bench Press (chest level)', category='bench', equipment='barbell', primary=['chest', 'triceps'], secondary=['frontDelts'], notes='Dead stop on pins set at chest level'),
    dict(id='bbm-deficit-sldl', name='Stiff-Leg Deadlift (2" deficit)', category='deadlift', equipment='barbell', primary=['hamstrings', 'glutes'], secondary=['lowerBack']),
    dict(id='gpp-upper-back', name='GPP: Upper back work', category='back', equipment='other', primary=['upperBack', 'lats'], secondary=['biceps'], notes='Any row, chin-up or pulldown variation'),
    dict(id='gpp-abs', name='GPP: Ab work', category='core', equipment='other', primary=['abs'], secondary=['obliques'], notes='Planks, ab wheel, leg raises, back extensions…'),
    dict(id='gpp-arms', name='GPP: Arm superset', category='arms', equipment='other', primary=['biceps', 'triceps'], secondary=[]),
    dict(id='gpp-cardio', name='GPP: Conditioning', category='conditioning', equipment='other', primary=[], secondary=[], notes='Log minutes as reps'),
]
for c in CUSTOM:
    c['custom'] = True

def rest_of(txt):
    t = (txt or '').lower()
    if '3-5' in t: return 240
    if '2-3' in t: return 150
    return 180

NUM = r'(\d+(?:\.\d+)?)'

def parse(text):
    """Returns (sets, short_notes) for one prescription cell."""
    sets, notes = [], []
    lines = []
    for raw in str(text).split('\n'):
        for part in re.split(r',\s*(?=\d+\s*reps?\s*@)', raw):  # "1 rep @ 8, 3 reps @ 85% e1RM"
            lines.append(part)
    last_rpe_set = None
    for raw in lines:
        l = raw.strip().lstrip('•').strip()
        low = l.lower()
        if not l or low in ('none',) or low.startswith('no back off'):
            continue
        # Load drop / take-off back-offs.
        if ('% from' in low or '% off' in low or low.startswith('take') or low.startswith('then take')) and '%' in low:
            pct = float(re.search(NUM + r'\s*%', l).group(1))
            m = re.search(r'sets of (\d+)', low) or re.search(r'x (\d+) reps', low) or re.search(r'do (\d+) sets? of (\d+)', low)
            reps = int(m.group(m.lastindex)) if m else (last_rpe_set['reps'] if last_rpe_set else 5)
            m = re.search(r'cap at (\d+)', low) or re.search(r'(\d+)-(\d+) (?:more )?sets', low) or re.search(r'for (\d+) more sets?', low) or re.search(r'x (\d+) sets?\b', low) or re.search(r'do (\d+) sets?', low)
            count = int(m.group(m.lastindex)) if m else 2
            if last_rpe_set: last_rpe_set['kind'] = 'top'
            sets += [dict(kind='backoff', reps=reps, percentOfTop=round(100 - pct, 1)) for _ in range(count)]
            if 'until' in low:
                tgt = re.search(r'until (?:effort is ~ )?@?\s*(?:rpe\s*)?' + NUM, low)
                notes.append(f'Back-offs: −{pct:g}% until RPE {tgt.group(1) if tgt else "target"} again (max {count})')
            if 'tempo' in low:
                tm = re.search(r'(\d-\d-\d)', l)
                notes.append(f'Back-offs as {tm.group(1) if tm else ""} tempo variation')
            continue
        # Percentage volume work.
        m = re.match(NUM + r'\s*%\s*e1rm\s*x\s*(\d+|amrap)\s*(?:reps?)?\s*(?:x\s*(\d+)\s*sets?)?(?:\s*of\s*(\d+)\s*reps)?(.*)', low)
        if m:
            p, r, k, of, rest = m.groups()
            if r == 'amrap':
                sets += [dict(kind='amrap', reps=5, percent=float(p)) for _ in range(int(k or 1))]
                notes.append(f'{p}% × AMRAP')
                continue
            r = int(r)
            sm = re.search(r'sets of (\d+)-(\d+)', rest)
            if not k and sm:  # total reps in sets of a-b
                per = int(sm.group(2)); k = math.ceil(r / per); r = per
                cap = re.search(r'rpe\s*(\d+(?:-\d+)?)', rest)
                notes.append(f'{m.group(2)} total reps at {p}%' + (f', keep sets ≤ RPE {cap.group(1)}' if cap else ''))
            sets += [dict(kind='work', reps=int(of or r), percent=float(p)) for _ in range(int(k or 1))]
            continue
        m = re.match(r'(\d+)\s*reps?\s*x\s*(\d+)\s*sets?\s*@\s*' + NUM + r'\s*%', low)
        if m:
            sets += [dict(kind='work', reps=int(m.group(1)), percent=float(m.group(3))) for _ in range(int(m.group(2)))]
            continue
        m = re.match(r'(\d+)\s*reps?\s*@\s*' + NUM + r'\s*%', low)
        if m:
            sets.append(dict(kind='work', reps=int(m.group(1)), percent=float(m.group(2))))
            continue
        # Repeat sets at the same load.
        m = re.match(r'repeat (\d+) @ ' + NUM + r' x (\d+) sets? of (\d+)', low)
        if m:
            if last_rpe_set: last_rpe_set['kind'] = 'top'
            sets += [dict(kind='backoff', reps=int(m.group(4)), percentOfTop=100) for _ in range(int(m.group(3)))]
            continue
        if 'opener' in low:
            sets.append(dict(kind='top', reps=1, percent=91)); notes.append('Opener (~91%): fast, confident single'); continue
        if low.startswith('work'):
            sets.append(dict(kind='work', reps=1, rpe=7)); notes.append('Work up to 1 rep at your last warm-up weight'); continue
        # RPE sets, optionally with an inline "Repeat" or "x N set".
        m = re.match(r'(\d+)\s*(?:reps?)?\s*@\s*(?:rpe\s*)?' + NUM, low)
        if m:
            s = dict(kind='work', reps=int(m.group(1)), rpe=float(m.group(2)))
            xs = re.search(r'x\s*(\d+)\s*set', low[m.end():]) if 'repeat' not in low else None
            for _ in range(int(xs.group(1)) if xs else 1):
                sets.append(dict(s))
            last_rpe_set = sets[-1]
            rp = re.search(r'repeat (\d+) @ ' + NUM + r' x (\d+) sets? of (\d+)', low)
            if rp:
                last_rpe_set['kind'] = 'top'
                sets += [dict(kind='backoff', reps=int(rp.group(4)), percentOfTop=100) for _ in range(int(rp.group(3)))]
            continue
        raise ValueError(f'Unparsed prescription line: {l!r}')
    for s in sets:
        if s.get('rpe') is not None and s['rpe'] == int(s['rpe']): s['rpe'] = int(s['rpe'])
    return sets, notes

def gpp_exercises(texts, freq_day):
    """GPP notes → exercises. freq_day=2 means the second lifting day that gets GPP (1x/wk items go there)."""
    out = []
    for t in texts:
        for line in str(t or '').split('\n'):
            l = line.strip(); low = l.lower()
            if not l or low == 'none': continue
            if '1x/wk' in low and freq_day == 1: continue
            if 'triceps' in low or 'biceps' in low:
                ex = 'triceps-pushdown' if 'triceps' in low else 'dumbbell-curl'
            elif 'arm' in low: ex = 'gpp-arms'
            elif 'upper back' in low: ex = 'gpp-upper-back'
            elif 'ab ' in low or low.startswith('ab') or ' ab' in low: ex = 'gpp-abs'
            elif 'steady state' in low or 'sprint' in low: ex = 'gpp-cardio'
            else: raise ValueError('GPP: ' + l)
            m = re.search(r'(\d+) sets? (?:of )?(?:upper back work x )?(\d+)-(\d+) reps @ rpe (\d+)', low)
            mins = re.search(r'(\d+) (?:min|minutes?)\b', low)
            if m:
                sets = [dict(kind='work', reps=int(m.group(2)), rpe=int(m.group(4))) for _ in range(int(m.group(1)))]
            elif mins:
                sets = [dict(kind='amrap' if 'amrap' in low else 'work', reps=int(mins.group(1)))]
            else:
                sets = [dict(kind='work', reps=10)]
            # Keep only what the set targets can't show.
            if m:
                rpe = re.search(r'rpe (\d+(?:-\d+)?)', low).group(1)
                note = f'{m.group(2)}-{m.group(3)} reps' + (f' @ RPE {rpe}' if '-' in rpe else '')
            elif ex == 'gpp-cardio':
                note = re.sub(r',?\s*\d+x/wk', '', l).strip() + ' · log minutes'
            elif mins:
                note = f'{mins.group(1)} min AMRAP · log minutes'
            else:
                note = ''
            out.append(dict(exerciseId=ex, sets=sets, rest=90, **({'notes': note} if note else {})))
    return out

weeks = []
for w in range(1, 14):
    ws = wb[f'WEEK {w}']
    block = ws['D3'].value or ''
    gpp = [ws['E121'].value, ws['K121'].value, ws['Q121'].value, ws['E125'].value]
    days = []
    for di, base in enumerate((26, 50, 74, 98)):
        exs = []
        meet_order = ['squat', 'bench-press', 'deadlift']
        for ci, col in enumerate((4, 10, 16)):
            name = ws.cell(base, col).value
            text = ws.cell(base + 1, col + 1).value
            if not name or str(name).strip().lower() == 'none' or not text or str(text).strip().lower() == 'none':
                continue
            name = str(name).strip()
            ex_id = meet_order[ci] if name == '#REF!' else NAME_MAP[name.lower()]
            sets, notes = parse(text)
            if not sets: continue
            if ex_id == 'pin-squat' : notes.insert(0, 'Pins just below parallel')
            if ex_id == 'deficit-deadlift': notes.insert(0, '2" deficit')
            exs.append(dict(exerciseId=ex_id, sets=sets, rest=rest_of(ws.cell(base + 3, col + 1).value), **({'notes': ' · '.join(notes)} if notes else {})))
        if not exs: continue
        if w <= 12 and di in (1, 3):
            exs += gpp_exercises(gpp, 1 if di == 1 else 2)
        name = f'Day {di + 1}'
        if w == 13 and di == 3: name = 'Meet / max test'
        days.append(dict(id=f'bbm-s3-w{w}-d{di + 1}', name=name, exercises=exs))
    weeks.append(dict(name=f'Week {w} · {block}', days=days))

program = dict(
    id='bbm-strength-iii',
    name='Strength III (Barbell Medicine)',
    author='Barbell Medicine',
    level='advanced',
    description='13-week strength block for powerlifting: 4 lifting days a week with RPE ramps, load drops and GPP, finishing with a taper and meet/max test. Imported from your template.',
    weeks=weeks,
)
json.dump(dict(format='strength-tracker-program', version=1, program=program, customExercises=CUSTOM), open(OUT, 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
n_days = sum(len(w['days']) for w in weeks); n_ex = sum(len(d['exercises']) for w in weeks for d in w['days'])
print(f'{len(weeks)} weeks, {n_days} days, {n_ex} exercise slots')
